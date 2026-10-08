"""Tabular adapters with bounded reads and fail-closed capability checks."""

from __future__ import annotations

import csv
import io
import inspect
import itertools
import json
import math
import warnings
import zipfile
from pathlib import Path
from typing import Any, Iterable, Sequence

from .common import (
    base_profile,
    key_diagnostics,
    survey_design,
    unsupported_profile,
    variable_record,
)
from .archive import ZipBoundaryError, zip_central_directory_summary
from .spss import profile_sav


MAX_JSON_PARSE_BYTES = 8 * 1024 * 1024
MAX_CSV_FIELD_BYTES = 1024 * 1024
MAX_CSV_SAMPLE_BYTES = 8 * 1024 * 1024
MAX_DTA_FILE_BYTES = 8 * 1024 * 1024 * 1024
MAX_EXCEL_FILE_BYTES = 512 * 1024 * 1024
MAX_EXCEL_ARCHIVE_MEMBERS = 10_000
MAX_EXCEL_CENTRAL_DIRECTORY_BYTES = 16 * 1024 * 1024
MAX_EXCEL_UNCOMPRESSED_BYTES = 256 * 1024 * 1024
MAX_EXCEL_MEMBER_UNCOMPRESSED_BYTES = 64 * 1024 * 1024
MAX_EXCEL_COMPRESSION_RATIO = 200.0
MAX_PROFILE_COLUMNS = 10_000
MAX_PROFILE_CELLS = 2_000_000
MAX_METADATA_TEXT_CHARS = 4096
MAX_HEADER_TEXT_BYTES = 2 * 1024 * 1024
HEADER_CONFIRMATION_REQUIRED = "HEADER_CONFIRMATION_REQUIRED"


class _ProfileResourceLimit(ValueError):
    """A resource boundary failure whose diagnostic never reflects input."""


class _CappedRawReader(io.RawIOBase):
    """Expose at most ``limit + 1`` bytes to a text decoder."""

    def __init__(self, handle: Any, limit: int) -> None:
        self._handle = handle
        self._remaining = limit + 1
        self.observed = 0
        self.limit = limit

    @property
    def exceeded(self) -> bool:
        return self.observed > self.limit

    def readable(self) -> bool:
        return True

    def readinto(self, buffer: bytearray) -> int:
        if self._remaining <= 0:
            return 0
        amount = min(len(buffer), self._remaining)
        data = self._handle.read(amount)
        if not data:
            return 0
        length = len(data)
        buffer[:length] = data
        self._remaining -= length
        self.observed += length
        return length


def _resource_block(
    profile: dict[str, Any],
    *,
    message: str,
    details: dict[str, int] | None = None,
) -> None:
    profile["parse_status"] = "BLOCKED"
    row: dict[str, Any] = {"code": "RESOURCE_LIMIT_EXCEEDED", "message": message}
    if details:
        row["details"] = details
    profile["errors"].append(row)


def _validated_names(values: Iterable[Any]) -> list[str]:
    names: list[str] = []
    counts: dict[str, int] = {}
    encoded_total = 0
    for index, value in enumerate(values):
        if index >= MAX_PROFILE_COLUMNS:
            raise _ProfileResourceLimit("column count exceeds the metadata policy limit")
        base = str(value).strip() if value not in (None, "") else f"column_{index + 1}"
        counts[base] = counts.get(base, 0) + 1
        name = base if counts[base] == 1 else f"{base}__{counts[base]}"
        if len(name) > MAX_METADATA_TEXT_CHARS:
            raise _ProfileResourceLimit("one or more metadata fields exceed the text limit")
        encoded_total += len(name.encode("utf-8"))
        if encoded_total > MAX_HEADER_TEXT_BYTES:
            raise _ProfileResourceLimit("header metadata exceeds the byte limit")
        names.append(name)
    return names


def validate_delimited_width(
    path: Path,
    *,
    delimiter: str,
    max_columns: int,
    max_bytes: int,
    record_limit: int,
) -> None:
    """Reject ultra-wide delimited records before ``csv`` allocates a row."""

    delimiter_bytes = delimiter.encode("ascii")
    if len(delimiter_bytes) != 1 or max_columns < 1 or record_limit < 1:
        raise _ProfileResourceLimit("invalid delimited parser boundary")
    separator = delimiter_bytes[0]
    in_quotes = False
    after_quote = False
    at_field_start = True
    record_nonempty = False
    columns = 1
    records = 0
    observed = 0
    first_chunk = True
    with path.open("rb") as handle:
        while records < record_limit:
            chunk = handle.read(min(64 * 1024, max_bytes + 1 - observed))
            if not chunk:
                break
            observed += len(chunk)
            if observed > max_bytes:
                raise _ProfileResourceLimit("delimited sample exceeds the byte limit")
            scan_chunk = chunk
            if first_chunk:
                first_chunk = False
                if scan_chunk.startswith(b"\xef\xbb\xbf"):
                    scan_chunk = scan_chunk[3:]
            for byte in scan_chunk:
                if in_quotes:
                    if byte == 0x22:
                        in_quotes = False
                        after_quote = True
                    continue
                if after_quote:
                    if byte == 0x22:
                        in_quotes = True
                        after_quote = False
                    elif byte == separator:
                        record_nonempty = True
                        columns += 1
                        after_quote = False
                        at_field_start = True
                        if columns > max_columns:
                            raise _ProfileResourceLimit(
                                "one or more delimited records exceed the column limit"
                            )
                    elif byte == 0x0A:
                        if record_nonempty:
                            records += 1
                        columns = 1
                        after_quote = False
                        at_field_start = True
                        record_nonempty = False
                        if records >= record_limit:
                            break
                    elif byte != 0x0D:
                        raise _ProfileResourceLimit(
                            "delimited sample violates the quoted-field policy"
                        )
                    continue
                if byte == 0x22 and at_field_start:
                    in_quotes = True
                    at_field_start = False
                    record_nonempty = True
                elif byte == separator:
                    record_nonempty = True
                    columns += 1
                    at_field_start = True
                    if columns > max_columns:
                        raise _ProfileResourceLimit(
                            "one or more delimited records exceed the column limit"
                        )
                elif byte == 0x0A:
                    if record_nonempty:
                        records += 1
                    columns = 1
                    at_field_start = True
                    record_nonempty = False
                    if records >= record_limit:
                        break
                elif byte != 0x0D:
                    at_field_start = False
                    record_nonempty = True
            if observed == max_bytes:
                if handle.read(1):
                    raise _ProfileResourceLimit("delimited sample exceeds the byte limit")
                break


def _check_cell_budget(row_limit: int, column_count: int) -> None:
    if row_limit > 0 and column_count > 0 and row_limit * column_count > MAX_PROFILE_CELLS:
        raise _ProfileResourceLimit("bounded row and column selection exceeds the cell limit")


def _safe_parse_failure(profile: dict[str, Any], *, code: str, message: str) -> None:
    profile["parse_status"] = "FAIL"
    profile["variables"] = []
    profile["column_count"] = 0
    profile["observed_rows"] = 0
    profile["errors"].append({"code": code, "message": message})


def _unconfirmed_header_profile(
    path: Path,
    *,
    adapter: str,
    backend: str,
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    """Block header-dependent formats without opening their contents.

    Delimited and spreadsheet formats do not carry a trustworthy machine-level
    distinction between a header row and an ordinary data row.  Treating the
    first row as metadata can therefore serialize private values as variable
    names.  Explicit confirmation is required before either the header or any
    bounded data rows are read.
    """

    profile = base_profile(path, adapter=adapter, backend=backend)
    profile["parse_status"] = "BLOCKED"
    profile["profile_scope"] = "header-confirmation"
    profile["capabilities"] = {
        "bounded_row_read": True,
        "header_confirmation_required": True,
        "header_confirmed": False,
        "full_file_fallback_allowed": False,
    }
    profile["errors"].append(
        {
            "code": HEADER_CONFIRMATION_REQUIRED,
            "message": (
                "explicit header confirmation is required before this format can be "
                "profiled; file contents were not read"
            ),
        }
    )
    return profile, []


def _dedupe_names(values: Iterable[Any]) -> list[str]:
    names: list[str] = []
    counts: dict[str, int] = {}
    for index, value in enumerate(values):
        base = str(value).strip() if value not in (None, "") else f"column_{index + 1}"
        counts[base] = counts.get(base, 0) + 1
        names.append(base if counts[base] == 1 else f"{base}__{counts[base]}")
    return names


def _mapping_rows(names: Sequence[str], values: Iterable[Sequence[Any]]) -> list[dict[str, Any]]:
    return [
        {name: row[index] if index < len(row) else None for index, name in enumerate(names)}
        for row in values
    ]


def _profile_stata(path: Path, row_limit: int) -> tuple[dict[str, Any], Any | None]:
    try:
        from pandas.io.stata import StataReader
    except ImportError:
        return (
            unsupported_profile(
                path,
                format_name="Stata DTA",
                dependency_requirement="the pinned pandas StataReader runtime",
            ),
            None,
        )

    try:
        supports_bounded_read = "nrows" in inspect.signature(StataReader.read).parameters
    except (TypeError, ValueError):
        supports_bounded_read = False
    if not supports_bounded_read:
        return (
            unsupported_profile(
                path,
                format_name="Stata DTA",
                dependency_requirement="a reader with a verified nrows-bounded API",
            ),
            None,
        )

    profile = base_profile(path, adapter="stata/v1", backend="pandas-stata-reader-bounded")
    profile["capabilities"] = {
        "dictionary_metadata": True,
        "bounded_row_read": True,
        "bounded_file_bytes": MAX_DTA_FILE_BYTES,
        "bounded_columns": MAX_PROFILE_COLUMNS,
        "bounded_cells": MAX_PROFILE_CELLS,
        "value_label_tables_read": False,
        "full_file_fallback_allowed": False,
    }
    frame = None
    try:
        if path.stat().st_size > MAX_DTA_FILE_BYTES:
            raise _ProfileResourceLimit("DTA file exceeds the parser byte limit")
        with warnings.catch_warnings(record=True) as reader_warnings:
            warnings.simplefilter("always")
            with StataReader(path, convert_categoricals=False, convert_missing=False) as reader:
                raw_names = list(getattr(reader, "_varlist", ()) or ())
                if len(raw_names) > MAX_PROFILE_COLUMNS:
                    raise _ProfileResourceLimit("DTA variable count exceeds the metadata policy limit")
                labels = reader.variable_labels()
                names = _validated_names(raw_names or labels)
                storage = list(getattr(reader, "_typlist", ()) or ())
                if len(labels) > MAX_PROFILE_COLUMNS or len(storage) > MAX_PROFILE_COLUMNS:
                    raise _ProfileResourceLimit("DTA dictionary metadata exceeds the column limit")
                for name in names:
                    label = str(labels.get(name) or "")
                    if len(label) > MAX_METADATA_TEXT_CHARS:
                        raise _ProfileResourceLimit("DTA variable labels exceed the metadata text limit")
                _check_cell_budget(row_limit, len(names))
                profile["file_row_count"] = int(getattr(reader, "_nobs", 0))
                data_label = getattr(reader, "data_label", "")
                profile["data_label"] = str(
                    data_label() if callable(data_label) else data_label or ""
                )
                if len(profile["data_label"]) > MAX_METADATA_TEXT_CHARS:
                    raise _ProfileResourceLimit("DTA data label exceeds the metadata text limit")
                profile["variables"] = [
                    variable_record(
                        str(name),
                        label=str(labels.get(name) or ""),
                        storage_type=str(storage[index]) if index < len(storage) else "unknown",
                    )
                    for index, name in enumerate(names)
                ]
                profile["column_count"] = len(profile["variables"])
                if row_limit > 0:
                    # The bounded API is a hard requirement. There is deliberately
                    # no fallback to pandas.read_stata or read() without nrows.
                    frame = reader.read(nrows=row_limit)
                    profile["observed_rows"] = int(len(frame))
                    profile["profile_scope"] = "bounded-rows"
                else:
                    profile["profile_scope"] = "dictionary-metadata"
        if reader_warnings:
            profile["errors"].append({
                "code": "DTA_READER_WARNING",
                "message": "dictionary metadata reader emitted a warning; decoded labels require verification",
            })
    except _ProfileResourceLimit:
        _resource_block(
            profile,
            message="DTA metadata exceeds the bounded parser resource policy",
        )
    except (OSError, ValueError, TypeError, NotImplementedError, MemoryError, RecursionError):
        _safe_parse_failure(
            profile,
            code="DTA_PARSE_FAILED",
            message="DTA metadata could not be parsed safely",
        )
    return profile, frame


def _profile_excel(
    path: Path,
    row_limit: int,
    *,
    header_confirmed: bool = False,
) -> tuple[dict[str, Any], Any | None]:
    if not header_confirmed:
        return _unconfirmed_header_profile(
            path,
            adapter="excel/v1",
            backend="openpyxl-read-only",
        )
    try:
        from openpyxl import load_workbook
    except ImportError:
        return (
            unsupported_profile(
                path,
                format_name="Excel XLSX",
                dependency_requirement="the pinned openpyxl read-only runtime",
            ),
            None,
        )

    profile = base_profile(path, adapter="excel/v1", backend="openpyxl-read-only")
    profile["capabilities"] = {
        "read_only_workbook": True,
        "bounded_row_read": True,
        "bounded_file_bytes": MAX_EXCEL_FILE_BYTES,
        "bounded_archive_members": MAX_EXCEL_ARCHIVE_MEMBERS,
        "bounded_central_directory_bytes": MAX_EXCEL_CENTRAL_DIRECTORY_BYTES,
        "bounded_uncompressed_bytes": MAX_EXCEL_UNCOMPRESSED_BYTES,
        "bounded_columns": MAX_PROFILE_COLUMNS,
        "bounded_cells": MAX_PROFILE_CELLS,
        "header_confirmation_required": True,
        "header_confirmed": True,
        "full_file_fallback_allowed": False,
    }
    workbook: Any | None = None
    rows_out: list[dict[str, Any]] = []
    try:
        if path.stat().st_size > MAX_EXCEL_FILE_BYTES:
            raise _ProfileResourceLimit("Excel file exceeds the parser byte limit")
        try:
            zip_central_directory_summary(
                path,
                member_limit=MAX_EXCEL_ARCHIVE_MEMBERS,
                central_directory_limit=MAX_EXCEL_CENTRAL_DIRECTORY_BYTES,
            )
        except ZipBoundaryError as exc:
            if exc.blocked:
                raise _ProfileResourceLimit("Excel archive metadata exceeds the parser policy") from None
            raise ValueError("invalid Excel ZIP metadata") from None
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_EXCEL_ARCHIVE_MEMBERS:
                raise _ProfileResourceLimit("Excel archive member count exceeds the parser policy")
            total_uncompressed = sum(int(item.file_size) for item in infos)
            largest_member = max((int(item.file_size) for item in infos), default=0)
            total_compressed = sum(int(item.compress_size) for item in infos)
            ratio = (
                math.inf
                if total_uncompressed > 0 and total_compressed <= 0
                else float(total_uncompressed) / max(1, total_compressed)
            )
            excessive_member_ratio = any(
                int(item.file_size) > 0
                and (
                    int(item.compress_size) <= 0
                    or float(item.file_size) / float(item.compress_size)
                    > MAX_EXCEL_COMPRESSION_RATIO
                )
                for item in infos
            )
            if (
                total_uncompressed > MAX_EXCEL_UNCOMPRESSED_BYTES
                or largest_member > MAX_EXCEL_MEMBER_UNCOMPRESSED_BYTES
                or ratio > MAX_EXCEL_COMPRESSION_RATIO
                or excessive_member_ratio
            ):
                raise _ProfileResourceLimit("Excel XML expansion exceeds the parser policy")
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        declared_columns = int(sheet.max_column or 0)
        if declared_columns > MAX_PROFILE_COLUMNS:
            raise _ProfileResourceLimit("Excel column count exceeds the metadata policy limit")
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, ())
        names = _validated_names(header)
        _check_cell_budget(row_limit, len(names))
        values = list(itertools.islice(rows, row_limit))
        rows_out = _mapping_rows(names, values)
        profile["variables"] = [variable_record(name) for name in names]
        profile["column_count"] = len(names)
        profile["file_row_count"] = max(0, int(sheet.max_row or 0) - 1)
        profile["observed_rows"] = len(rows_out)
        profile["profile_scope"] = "bounded-rows" if row_limit else "dictionary-metadata"
    except _ProfileResourceLimit:
        rows_out = []
        _resource_block(
            profile,
            message="Excel metadata exceeds the bounded parser resource policy",
        )
    except (
        OSError,
        ValueError,
        TypeError,
        IndexError,
        KeyError,
        zipfile.BadZipFile,
        MemoryError,
        RecursionError,
    ):
        rows_out = []
        _safe_parse_failure(
            profile,
            code="EXCEL_PARSE_FAILED",
            message="Excel metadata could not be parsed safely",
        )
    finally:
        if workbook is not None:
            workbook.close()
    return profile, rows_out


def _profile_csv(
    path: Path,
    row_limit: int,
    *,
    header_confirmed: bool = False,
) -> tuple[dict[str, Any], Any | None]:
    if not header_confirmed:
        return _unconfirmed_header_profile(
            path,
            adapter="delimited/v1",
            backend="stdlib-csv-bounded",
        )
    profile = base_profile(path, adapter="delimited/v1", backend="stdlib-csv-bounded")
    profile["capabilities"] = {
        "bounded_row_read": True,
        "field_size_limit_bytes": MAX_CSV_FIELD_BYTES,
        "sample_read_limit_bytes": MAX_CSV_SAMPLE_BYTES,
        "bounded_columns": MAX_PROFILE_COLUMNS,
        "bounded_cells": MAX_PROFILE_CELLS,
        "header_confirmation_required": True,
        "header_confirmed": True,
        "full_file_fallback_allowed": False,
    }
    rows_out: list[dict[str, Any]] = []
    prior_field_limit = csv.field_size_limit()
    try:
        csv.field_size_limit(MAX_CSV_FIELD_BYTES)
        delimiter = "\t" if path.suffix.casefold() == ".tsv" else ","
        validate_delimited_width(
            path,
            delimiter=delimiter,
            max_columns=MAX_PROFILE_COLUMNS,
            max_bytes=MAX_CSV_SAMPLE_BYTES,
            record_limit=max(1, row_limit + 1),
        )
        with path.open("rb") as binary_handle:
            capped = _CappedRawReader(binary_handle, MAX_CSV_SAMPLE_BYTES)
            with io.BufferedReader(capped) as buffered, io.TextIOWrapper(
                buffered,
                encoding="utf-8-sig",
                newline="",
                errors="strict",
            ) as text_handle:
                reader = csv.reader(text_handle, delimiter=delimiter, strict=True)
                names = _validated_names(next(reader, ()))
                _check_cell_budget(row_limit, len(names))
                values = []
                for row in itertools.islice(reader, row_limit):
                    if len(row) > MAX_PROFILE_COLUMNS:
                        raise _ProfileResourceLimit(
                            "one or more delimited rows exceed the column limit"
                        )
                    values.append(row)
                if capped.exceeded:
                    raise _ProfileResourceLimit(
                        "bounded delimited rows exceed the sample byte limit"
                    )
        rows_out = _mapping_rows(names, values)
        profile["variables"] = [variable_record(name) for name in names]
        profile["column_count"] = len(names)
        profile["observed_rows"] = len(rows_out)
        profile["profile_scope"] = "bounded-rows" if row_limit else "dictionary-metadata"
    except _ProfileResourceLimit:
        rows_out = []
        _resource_block(
            profile,
            message="Delimited metadata exceeds the bounded parser resource policy",
        )
    except (OSError, UnicodeError, csv.Error, ValueError, MemoryError, RecursionError):
        rows_out = []
        _safe_parse_failure(
            profile,
            code="DELIMITED_PARSE_FAILED",
            message="Delimited metadata could not be parsed safely",
        )
    finally:
        csv.field_size_limit(prior_field_limit)
    return profile, rows_out


def _profile_json(path: Path, row_limit: int) -> tuple[dict[str, Any], Any | None]:
    profile = base_profile(path, adapter="json/v1", backend="stdlib-json-size-capped")
    profile["capabilities"] = {
        "bounded_file_bytes": MAX_JSON_PARSE_BYTES,
        "bounded_output_rows": row_limit,
        "bounded_columns": MAX_PROFILE_COLUMNS,
        "bounded_cells": MAX_PROFILE_CELLS,
        "full_file_fallback_allowed": False,
    }
    rows_out: list[dict[str, Any]] = []
    try:
        with path.open("rb") as handle:
            raw = handle.read(MAX_JSON_PARSE_BYTES + 1)
        if len(raw) > MAX_JSON_PARSE_BYTES:
            profile["parse_status"] = "BLOCKED"
            profile["profile_scope"] = "size-cap"
            profile["errors"].append(
                {
                    "code": "RESOURCE_LIMIT_EXCEEDED",
                    "message": (
                        "JSON exceeds the bounded parser byte limit; no full-file fallback was attempted"
                    ),
                    "details": {
                        "observed_at_least_bytes": len(raw),
                        "allowed_bytes": MAX_JSON_PARSE_BYTES,
                    },
                }
            )
            return profile, rows_out
        try:
            payload = json.loads(raw.decode("utf-8-sig"))
        except RecursionError:
            raise _ProfileResourceLimit("JSON nesting exceeds the parser policy") from None
        candidates = payload if isinstance(payload, list) else [payload]
        rows_out = [row for row in candidates[:row_limit] if isinstance(row, dict)]
        dictionary_rows = rows_out
        if row_limit == 0:
            dictionary_rows = [
                row for row in candidates[:100] if isinstance(row, dict)
            ][:1]
        unique_keys: set[str] = set()
        for row in dictionary_rows:
            for key in row:
                unique_keys.add(str(key))
                if len(unique_keys) > MAX_PROFILE_COLUMNS:
                    raise _ProfileResourceLimit("JSON object keys exceed the column limit")
        names = _validated_names(sorted(unique_keys))
        _check_cell_budget(row_limit, len(names))
        profile["variables"] = [variable_record(name) for name in names]
        profile["column_count"] = len(names)
        profile["file_row_count"] = len(payload) if isinstance(payload, list) else 1
        profile["observed_rows"] = len(rows_out)
        profile["profile_scope"] = "bounded-rows" if row_limit else "dictionary-metadata"
    except _ProfileResourceLimit:
        rows_out = []
        _resource_block(
            profile,
            message="JSON metadata exceeds the bounded parser resource policy",
        )
    except (OSError, UnicodeError, json.JSONDecodeError, MemoryError, ValueError):
        rows_out = []
        _safe_parse_failure(
            profile,
            code="JSON_PARSE_FAILED",
            message="JSON metadata could not be parsed safely",
        )
    return profile, rows_out


def _profile_sas(path: Path) -> tuple[dict[str, Any], None]:
    return (
        unsupported_profile(
            path,
            format_name="SAS",
            dependency_requirement="a pinned parser with a verified bounded-read API",
        ),
        None,
    )


def profile_tabular(
    path: Path,
    *,
    row_limit: int,
    key_fields: Sequence[str],
    survey: bool,
    header_confirmed: bool = False,
) -> dict[str, Any]:
    suffix = path.suffix.casefold()
    if suffix == ".sav":
        profile = profile_sav(path)
        frame = None
    elif suffix == ".dta":
        profile, frame = _profile_stata(path, row_limit)
    elif suffix in {".xlsx", ".xlsm"}:
        profile, frame = _profile_excel(
            path,
            row_limit,
            header_confirmed=header_confirmed,
        )
    elif suffix in {".csv", ".tsv"}:
        profile, frame = _profile_csv(
            path,
            row_limit,
            header_confirmed=header_confirmed,
        )
    elif suffix == ".json":
        profile, frame = _profile_json(path, row_limit)
    elif suffix in {".sas7bdat", ".xpt"}:
        profile, frame = _profile_sas(path)
    else:
        profile = unsupported_profile(
            path,
            format_name=suffix.lstrip(".") or "extensionless binary",
            dependency_requirement="a registered, bounded adapter",
        )
        frame = None
    profile["key_diagnostics"] = key_diagnostics(frame, key_fields) if frame is not None else {
        "requested_fields": list(key_fields),
        "observed_fields": [],
        "missing_fields": list(key_fields),
        "sample_rows": 0,
        "nonmissing_key_rows": 0,
        "duplicate_key_rows": 0,
        "unique_in_bounded_sample": False,
        "claim_boundary": "no bounded rows were read",
    }
    profile["survey_design"] = survey_design(profile.get("variables") or [], expected=survey)
    return profile
